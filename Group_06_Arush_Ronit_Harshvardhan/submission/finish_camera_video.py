"""Build the shareable camera demo, analysis card and timed narration documents."""
from pathlib import Path
from zipfile import ZipFile,ZIP_STORED,ZIP_DEFLATED
import argparse,csv,hashlib,json,subprocess,textwrap
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]
def font(size):
    try:return ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',size)
    except OSError:return ImageFont.load_default(size=size)
def stamp(seconds):
    minutes=int(seconds//60);return f'{minutes:02d}:{int(seconds%60):02d}'
def probe(path):
    return json.loads(subprocess.check_output(['ffprobe','-v','error','-count_frames','-select_streams','v:0','-show_entries','stream=codec_name,width,height,r_frame_rate,nb_read_frames,duration','-show_entries','format=duration','-of','json',str(path)],text=True))


def build(run_dir,output_dir,analysis_seconds=12):
    run_dir=Path(run_dir);output_dir=Path(output_dir)
    report=json.loads((run_dir/'result.json').read_text());checks=json.loads((run_dir/'independent_collection_checks.json').read_text())
    assert checks['pass'] and report['status']=='CLEAN_AND_HOME','Only a verified complete mission can be published'
    with (run_dir/'frame_map.csv').open() as stream:frames=list(csv.DictReader(stream))
    raw_probe=probe(run_dir/'cinematic_cleaning.mp4');raw=raw_probe['streams'][0]
    assert len(frames)==int(raw['nb_read_frames'])
    assert all(int(row['frame'])==i and abs(float(row['video_time_s'])-i/25)<1e-6 for i,row in enumerate(frames))
    raw_duration=len(frames)/25
    output_dir.mkdir(parents=True,exist_ok=False)
    business=json.loads((ROOT/'analytics/final_business_v1/results.json').read_text());base=business['base_case']
    card=Image.new('RGB',(1280,720),(13,23,32));draw=ImageDraw.Draw(card)
    draw.text((45,20),'SCENARIO BUSINESS ANALYSIS',font=font(32),fill=(225,239,241))
    draw.text((45,66),'Simulation throughput + declared commercial assumptions / not measured field savings',font=font(18),fill=(180,209,215))
    chart=Image.open(ROOT/'analytics/final_business_v1/business_sensitivity.png').convert('RGB');chart.thumbnail((1190,440),Image.Resampling.LANCZOS)
    card.paste(chart,((1280-chart.width)//2,106))
    for x,value,label in [(45,f"{base['NCRI']*100:.1f}%",'Scenario normalised cost reduction'),(450,f"{base['payback_months']:.1f} months",'Scenario simple payback'),(855,f"{base['labour_reclaimed_person_h_month']:.1f} h/month",'Scenario labour reclaimed')]:
        draw.rounded_rectangle((x,558,x+380,651),radius=10,fill=(27,45,55))
        draw.text((x+14,570),value,font=font(29),fill=(225,215,170));draw.text((x+14,614),label,font=font(16),fill=(199,219,225))
    draw.text((45,675),'Manual bay transfer and maintenance included. No PV electrical yield uplift is monetised.',font=font(17),fill=(180,209,215))
    card_path=output_dir/'business_analysis_card.png';card.save(card_path)
    video=output_dir/'Group06_Complete_Simulation_With_Cameras.mp4'
    subprocess.run(['ffmpeg','-hide_banner','-loglevel','warning','-n','-i',str(run_dir/'cinematic_cleaning.mp4'),'-loop','1','-framerate','25','-i',str(card_path),'-filter_complex',f'[0:v]setpts=PTS-STARTPTS[v0];[1:v]trim=duration={analysis_seconds},setpts=PTS-STARTPTS,format=yuv420p[v1];[v0][v1]concat=n=2:v=1:a=0[v]','-map','[v]','-an','-r','25','-c:v','libx264','-preset','fast','-crf','23','-threads','2','-pix_fmt','yuv420p','-movflags','+faststart','-t',str(raw_duration+analysis_seconds),str(video)],check=True)
    final_probe=probe(video);duration=float(final_probe['format']['duration'])
    subprocess.run(['ffmpeg','-v','error','-i',str(video),'-f','null','-'],check=True)
    # Physical events are matched to the first sampled video state at/after their time.
    def event_time(simulation_time):
        for row in frames:
            if float(row['simulation_time_s'])>=simulation_time:return float(row['video_time_s'])
        return raw_duration-3
    solid_contact=event_time(report['physical_solid']['first_tool_contact_s'])
    tray_event=next(e['time_s'] for e in report['physical_solid']['events'] if e['event']=='SOLID_TRAY_CONTACT')
    tray_time=event_time(tray_event)
    arush_end=max(75.,min(90.,solid_contact-5.))
    ronit_end=min(duration-55.,max(arush_end+45.,tray_time+12.))
    paragraphs=[
      ('Arush',0.,arush_end,
       'My main contribution was the 3D robot model, controller coding and integration of this MuJoCo simulation.\n\n'
       'These close-ups show both panel inspection cameras and the onboard camera at their calibrated sensor positions. Their actual RGB inputs appear on the right, with green detection boxes.\n\n'
       'The robot begins at an external dock. Persistent dirt detections trigger its mission. I modelled mass, motor limits, rotating joints and collision geometry, then linked these to the controller.\n\n'
       'I used a back-and-forth coverage route with proportional-integral wheel-speed control. Motor torques are bounded, and the safety supervisor monitors chassis height, tilt and stopping behaviour. Camera observations can request local slowdowns and targeted revisits when dirt remains. Position feedback currently comes from the simulator.\n\n'
       'The brush is physically connected to the chassis. Fine-dirt patches shrink when the rotating brush makes qualifying contact, and removed mass is tracked through the onboard bin. This capture process is a reduced-order model; suction airflow is not resolved. The completed run clears all thirty-two selected patches below the threshold, with approximately ninety percent qualifying brush-contact coverage.'),
      ('Ronit',arush_end,ronit_end,
       'I helped refine the 3D model and contributed to the business modelling. The design connects an external docking platform, fixed brush supports, a solids guide and a permanent catch tray. Keeping the home platform off the glass also leaves that corner available for the cleaning route.\n\n'
       'Here, the larger object is a persistent rigid body with mass and gravity. The diagonal guide physically redirects it off the glass into the tray through collision contact. This test uses a prescribed handling corridor.\n\n'
       'The fixed mechanism keeps the concept relatively simple. I also helped connect practical operating requirements, including supervision, maintenance and panel-to-panel transfer, to the economic analysis.'),
      ('Harshvardhan',ronit_end,duration,
       'My focus was the business and economics of the product. I worked on capital and operating costs, labour requirements, payback and sensitivity analysis. The robot finishes its cycle by returning to the dock for modelled emptying and charging service.\n\n'
       'The operating model uses measured simulation throughput, while commercial inputs remain declared assumptions. It includes human supervision, maintenance and manual transfer between separate panel sections.\n\n'
       f"Under this scenario, the model estimates about {base['NCRI']*100:.0f} percent normalised cost reduction, a {base['payback_months']:.1f}-month simple payback and {base['labour_reclaimed_person_h_month']:.0f} person-hours reclaimed per month. These are scenario outputs, not measured commercial savings.\n\n"
       'The final graph shows how the conclusion changes with supervision and other uncertain inputs. Prototype traction, coating safety and suction performance remain future tests.')
    ]
    title='Group 06 — voice-over for the complete camera simulation'
    face_budget=max(0.,300-duration)
    lines=[f'# {title}',f'\nVideo length: **{stamp(duration)}** ({duration:.2f} seconds). The MP4 is silent so the three members can record their own voices. Intro/outro face clips are separate. For a five-minute final post, the combined face introduction and conclusion budget is approximately **{face_budget:.0f} seconds**.\n',
           'Speak at the rates shown below; short pauses can follow the camera and solid-object close-ups. These are suggested project-contribution statements based on the team roles supplied by the user.\n']
    doc=Document();doc.add_heading(title,0);doc.add_paragraph(f'Video length: {stamp(duration)}. Three-person narration. Record your own voices; no synthetic narration is included.')
    cues=[]
    for name,start,end,text in paragraphs:
        words=len(text.split());rate=words/(end-start)*60
        heading=f'{stamp(start)}–{stamp(end)} | {name}'
        lines += [f'## {heading}',f'\n{words} words; approximately {rate:.0f} words/minute.\n',text+'\n']
        doc.add_heading(heading,1);doc.add_paragraph(f'{words} words; approximately {rate:.0f} words/minute.')
        for paragraph in text.split('\n\n'):doc.add_paragraph(paragraph)
        cues.append({'speaker':name,'start_s':start,'end_s':end,'words':words,'words_per_minute':rate})
    lines += ['## On-screen editing cues',
              '- 00:00–00:08: home platform, falling dirt and the RGB trigger.',
              '- 00:08–00:25: wide camera rigs, left camera, right camera and onboard camera close-ups; actual RGB inputs stay visible.',
              f'- {stamp(solid_contact)}: first physical guide/object contact; {stamp(tray_time)}: object reaches the tray.',
              f'- {stamp(raw_duration-3)}: held measured cleaning result; {stamp(raw_duration)}: assumption-based business graph.',
              '- Traversal acceleration is labelled in the footage. This is recorded fresh simulation, not a continuously live feed.',
              '- Fine dust and dock service are reduced-order models; the larger block uses actual contact dynamics. Navigation uses simulator pose.']
    (output_dir/'VOICEOVER_SCRIPT.md').write_text('\n\n'.join(lines)+'\n')
    (output_dir/'VOICEOVER_SCRIPT.txt').write_text('\n\n'.join([title]+[f'{stamp(a)}–{stamp(b)} {name}\n\n{text}' for name,a,b,text in paragraphs])+'\n')
    doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(11);doc.save(output_dir/'VOICEOVER_SCRIPT.docx')
    (output_dir/'VOICEOVER_CUES.json').write_text(json.dumps(cues,indent=2)+'\n')
    timeline=[]
    for label,time in [('Off-glass home',0.),('Dirt/RGB trigger',4.),('Two camera rigs',8.),('Left panel camera',11.5),('Right panel camera',15.),('Onboard camera',18.5),('Camera-rig overview',22.)]:timeline.append({'feature':label,'video_time_s':event_time(time),'simulation_time_s':time})
    timeline += [{'feature':'Solid first tool contact','video_time_s':solid_contact,'simulation_time_s':report['physical_solid']['first_tool_contact_s']},{'feature':'Solid tray contact','video_time_s':tray_time,'simulation_time_s':tray_event},{'feature':'Assumption-based business analysis card','video_time_s':raw_duration,'simulation_time_s':None}]
    (output_dir/'VIDEO_TIMELINE.json').write_text(json.dumps(timeline,indent=2)+'\n')
    receipt={'pass':True,'mission_independently_verified':True,'video_seconds':duration,'frames':int(final_probe['streams'][0]['nb_read_frames']),'resolution':[1280,720],'full_decode_pass':True,'raw_capture_frames_have_provenance':True,'business_card_s':analysis_seconds,'audio':'None; record the supplied human voice-over','physical_source':str(run_dir.relative_to(ROOT)),'checks':str((run_dir/'independent_collection_checks.json').relative_to(ROOT)),'raw_video_sha256':hashlib.sha256((run_dir/'cinematic_cleaning.mp4').read_bytes()).hexdigest(),'final_video_sha256':hashlib.sha256(video.read_bytes()).hexdigest()}
    (output_dir/'VIDEO_CHECKS.json').write_text(json.dumps(receipt,indent=2)+'\n');(output_dir/'ffprobe.json').write_text(json.dumps(final_probe,indent=2)+'\n')
    for name in ['camera_rigs_overview.png','left_panel_camera.png','right_panel_camera.png','onboard_camera.png','camera_rigs_and_panel.png','solid_contact_closeup.png','solid_in_tray.png','dock_service.png','cinematic_final.png']:
        source=run_dir/name
        if source.exists():(output_dir/name).write_bytes(source.read_bytes())
    (output_dir/'README.md').write_text(f'# Complete simulation with visible cameras\n\nVideo: {stamp(duration)}, 720p/25 fps, silent for your own narration. This is a complete fresh actual-physics recording with camera hardware and true RGB feeds. No robot poses are animated or replayed into the simulation. A separately labelled {analysis_seconds}-second economic-analysis card follows the measured mission result.\n\nThe three-person script is provided as Word, plain text and Markdown with timing cues. Arush explains the primary model/code/simulation contribution; Ronit explains model refinement and business support; Harshvardhan explains economics. Review the supplied role statements before recording. Face introduction and conclusion clips are separate.\n\nThe underlying combined mission passes independent 500 Hz motion, camera causality, dirt, mass, home/service and physical solid retention checks. Selected dust is below its residual threshold; whole-panel coverage is approximately 90%, not 100%. Fine dust/service are reduced-order models and containment guards remain part of the physics. Commercial values are scenario assumptions, not field savings or measured PV power recovery. Raw telemetry, exact frame-to-simulation mapping, model and executed source snapshots remain in the linked run directory.\n')
    bundle=output_dir/'Group06_Camera_Demo_Package.zip'
    with ZipFile(bundle,'w',ZIP_DEFLATED) as archive:
        for path in sorted(output_dir.iterdir()):
            if path.is_file() and path!=bundle:archive.write(path,path.name,compress_type=ZIP_STORED if path.suffix in ('.mp4','.png','.docx') else ZIP_DEFLATED)
    with ZipFile(bundle) as archive:assert archive.testzip() is None
    assert video.stat().st_size<100*1024**2 and bundle.stat().st_size<100*1024**2
    print(json.dumps({'video':str(video),'duration_s':duration,'voiceover':cues,'package':str(bundle)},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--run-dir',type=Path,required=True);parser.add_argument('--output-dir',type=Path,required=True);parser.add_argument('--analysis-seconds',type=int,default=12);args=parser.parse_args();build(args.run_dir,args.output_dir,args.analysis_seconds)
