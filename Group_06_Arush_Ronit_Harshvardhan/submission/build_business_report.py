"""Build a standalone business report from the frozen final analysis.

No simulation outcomes or commercial observations are generated here. Tables
read retained results; the proposed commercial model is labelled as a proposal.
Use a new output directory to avoid overwriting any team-edited document.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "analytics/final_business_v1"


def read(path):
    return json.loads(path.read_text())


def build(output):
    output.mkdir(parents=True, exist_ok=False)
    report = read(DATA / "results.json")
    a, m, b = report["assumptions"], report["measured"], report["base_case"]
    conservative = report["success_lower_confidence_scenario"]
    sensitivity = pd.read_csv(DATA / "sensitivity.csv")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in [("Title", 25), ("Heading 1", 17), ("Heading 2", 12)]:
        document.styles[name].font.name = "Calibri"
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor.from_string("153C4A")
    header = section.header.paragraphs[0]
    header.text = "GROUP 06  |  SOLAR PANEL CLEANING ROBOT  |  BUSINESS MODEL"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("55717A")
    footer = section.footer.paragraphs[0]
    footer.text = "MDRIIA · 702CO0E012  |  Scenario economics, calibrated to simulation  |  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.runs[0].font.size = Pt(8)

    def p(text, bold=False):
        node = document.add_paragraph()
        node.add_run(text).bold = bold
        return node

    def h(text, level=1):
        document.add_heading(text, level=level)

    def page(title):
        document.add_page_break()
        h(title)

    def table(headers, rows, size=9, widths=None):
        t = document.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = str(text)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E5EFF1")
            cell._tc.get_or_add_tcPr().append(shade)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for ri, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(size)
                        run.bold = ri == 0
        if widths:
            t.autofit = False
            for column, width in zip(t.columns, widths):
                column.width = Inches(width)
            for row in t.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
        return t

    def picture(path, caption):
        document.add_picture(str(path), width=Inches(6.65))
        node = p(caption)
        node.runs[0].font.size = Pt(8)

    # 1. Executive summary.
    document.add_heading("Business Model &\nTechno-Economic Analysis", level=0)
    p("Autonomous inclined-solar-panel cleaning crawler", True)
    p("Group 06 · Arush Ashish Patil (E048), Ronit Rajput (E052), Harshvardhan Sahi (E058)")
    p("Business-model script and economics: Harshvardhan. Business/model support: Ronit. Robot physics, control logic and simulation: Arush. Development and report preparation were AI-assisted.")
    h("Decision summary", 2)
    p("The proposed product combines contamination-triggered cleaning, a fixed rotary brush, a solid-object guide/catch tray and an external service dock. The business case evaluates whether one cleaner serving manually transferred panel sections can reduce operating effort and cost relative to an assumed manual baseline.")
    table(["Base-case output", "Calculated result"], [
        ["Normalised annual operating cost, κ", f"{b['kappa']:.6f} of the manual baseline"],
        ["Normalised Cost Reduction Index, 1 − κ", f"{100*b['NCRI']:.2f}%"],
        ["Simple payback / lifetime-normalised payback", f"{b['payback_months']:.2f} months / {b['lifetime_normalised_payback']:.4f}"],
        ["Monthly human effort / reclaimed effort", f"{b['total_robot_related_person_h_month']:.2f} / {b['labour_reclaimed_person_h_month']:.2f} person-hours"],
        ["Service capacity / assumed demand", f"{b['service_rate_sections_h']:.3f} / {b['arrival_rate_sections_h']:.3f} sections/hour"],
        ["M/M/1 mean queue wait", f"{60*b['queue_wait_h']:.2f} minutes under distribution assumptions"],
    ])
    h("What these numbers establish", 2)
    p("They are reproducible scenario outputs using measured simulation timing/work and explicitly assumed commercial inputs. They are not supplier quotations, customer-trial savings, measured electrical consumption or PV energy-yield recovery. Currency values are excluded in accordance with the course rubric.")
    p("The report covers the product proposition, operating workflow, assumptions, equations, worked calculations, queueing, sensitivity, uncertainty, risk and reproduction. The companion package contains the actual Python business script and its input/output files.")

    # 2. Business proposition, clearly proposed rather than market-validated.
    page("1 · Proposed commercial model")
    p("This is a proposed route to market; customer demand, willingness to pay and supplier costs have not been validated.")
    table(["Business-model element", "Proposed approach"], [
        ["Customer segments", "Commercial rooftop-PV operators and operations/maintenance contractors with accessible, compatible panel sections."],
        ["Value proposition", "Clean when dirt is detected; reduce hands-on sweeping; combine a visible digital prototype with a transparent labour/cost model."],
        ["Differentiation", "RGB-gated departure, inspection feeds, targeted revisits, a fixed broad brush, rigid-body handling of a selected larger object and external servicing."],
        ["Offering / revenue hypothesis", "Equipment sale or lease with maintenance support. A managed cleaning service priced per section is an alternative; no tariff or profit is calculated here."],
        ["Customer relationships", "Site assessment, installation/rehearsal, operator training, periodic inspection and maintenance records."],
        ["Channels / partners", "Proposed direct pilot with a PV operator or cleaning contractor; prospective fabrication, motor, camera and service partners. No partnership is claimed."],
        ["Key resources / activities", "Crawler, brush/guide, cameras, containment, dock/tray, software, model validation, inspection and manual transfer between sections."],
        ["Cost structure", "Capital equipment plus supervision, transfer/recovery labour, maintenance, estimated energy, spares/battery and consumables."],
        ["Adoption gate", "Demonstrate real traction, panel/coating compatibility, safe mounting, reliable sensing, measured electrical power and operator workflow before commercial claims."],
    ], widths=[1.6, 5.35])
    h("Operating workflow and responsibility", 2)
    p("The operator places and checks the system on one modelled 6 × 4 m bay. Cameras monitor contamination; persistent detections trigger cleaning. The robot sweeps, revisits detected residuals, guides the selected solid to the tray, returns to the external dock and performs modelled emptying/charging service. A human transfers it to the next bay, handles recovery and performs maintenance.")
    p("The model is a cost-feasibility calculation, not a complete vendor profit-and-loss forecast. Selling price, customer acquisition, manufacturing scale, financing, tax and discounted cash flows require additional evidence.")

    # 3. Measured evidence and operational assumptions.
    page("2 · Simulation evidence and workload")
    p("Economic calibration selects all 30 proposed-controller attempts at 20° from the frozen 180-attempt evaluation (two controllers × three inclinations × 30 attempts). Failures are retained by the selector; all 30 selected attempts happened to succeed.")
    table(["Simulation-derived input", "Value", "Interpretation"], [
        ["Selected sample", f"{m['successes']}/{m['n']} completed", "Observed simulation fraction p = 1; not proof of perfect field reliability."],
        ["Mean attempt time Tₐ", f"{m['mean_attempt_s']:.6f} s", "About 10.21 minutes; full attempt timing, not video runtime."],
        ["Positive actuator mechanical work", f"{m['mean_positive_mechanical_work_J']:.6f} J", "Integrated simulated torque × speed; not electrical energy."],
        ["Modelled blower energy", f"{m['mean_modelled_blower_J']:.6f} J", "Computed using assumed 20 W blower power; not a power measurement."],
        ["Input provenance", "summary.csv + protocol/statistics", "analytics/final_validation_v1; raw 500 Hz trial logs retained separately."],
    ])
    h("Operational scenario assumptions", 2)
    table(["Input", "Assumed value"], [
        ["Section size / workload", "24 m²; 21 sections/day; 30 days/month; 8 scheduled hours/day"],
        ["Derived workload", "504 modelled m²/day; 630 sections/month; 15,120 modelled m²/month"],
        ["Manual baseline H₀", "120 person-hours/month for the specified workload"],
        ["Supervision fraction f", "0.20 of active attempt time"],
        ["Transfer to next section", "180 s/section, fully counted as human labour"],
        ["Recovery after a failed attempt", "180 s/failure, fully counted as human labour"],
        ["Maintenance", "10 person-hours/month"],
        ["Retries / queue", "Independent repeat attempts; one server; Poisson arrivals and exponential service assumed"],
    ])
    p("The workload is a declared operating scenario, not a measured installation size or inferred 100 kWp plant. Modelled surface area is not certified cleaned area. Autonomous transfer between separate bays has not been demonstrated.")

    # 4. Remaining financial/energy assumptions.
    page("3 · Cost and energy assumptions")
    p("Let one year of manual baseline labour cost equal 1 cost unit. All costs below are dimensionless ratios or normalised coefficients, not currency amounts. Commercial inputs are assumptions; no supplier quote is represented as evidence.")
    table(["Input", "Assumption / units", "Where used"], [
        ["Capital ratio k", "0.65 × one-year manual baseline cost", "Simple payback; separate from annual operating κ"],
        ["Equipment life L", "5 years", "Lifetime-normalised payback"],
        ["Battery / spares annual ratio", "0.05", "Annual operating-cost ratio"],
        ["Consumables annual ratio", "0.03", "Annual operating-cost ratio"],
        ["Electricity coefficient α", "0.0003 baseline cost units/kWh", "Annual estimated energy × α"],
        ["Electromechanical efficiency η", "0.60", "Convert positive simulated actuator work to an estimate"],
        ["Robot electronics", "5 W", "Attempts, recovery and transfer"],
        ["Panel cameras", "2 cameras × 3 W each", "8 scheduled hours/day"],
        ["Blower", "20 W during modelled blower activity", "Already included in modelled blower energy"],
        ["Human fractions", "Transfer = 1; failure recovery = 1", "No free inter-bay transfer or recovery labour"],
    ])
    h("Scope of the financial comparison", 2)
    p("κ compares recurring annual operations against one annual manual baseline. It does not include annualised capital expenditure. Capital is evaluated separately in the simple-payback equation. This avoids counting the same capital expense twice.")
    p("The model includes an assumed efficiency and auxiliary loads, but does not resolve motor-controller losses, battery ageing, standby demand or actual charging efficiency in hardware. Electrical figures are estimates. Dock energy credit in the simulation is not evidence of a full battery recharge.")
    h("Benefits deliberately not monetised", 2)
    p("PV electrical-yield recovery, measured water savings, coating life, quantified safety benefits and avoided downtime are not monetised. The cleaning model tracks contamination and contact, not PV power generation. No currency-based selling price, bill of materials or verified commercial profit is claimed.")

    # 5. Equations directly matching calculate().
    page("4 · Calculation model and labour")
    table(["Quantity", "Equation / definition"], [
        ["Monthly completed workload", "n = sections/day × operating days/month = 630"],
        ["Expected recovery per success", "R = (1 − p) × tᵣ / p"],
        ["Time per successful service", "Tₛ = Tₐ / p + R + tₜ; tₜ is transfer time"],
        ["Active attempt hours/month", "Hₐ = n × Tₐ / (3600p)"],
        ["Supervision hours", "Hsupervision = f × Hₐ"],
        ["Transfer / recovery hours", "Htransfer = n × (tₜ × hₜ + R × hᵣ) / 3600"],
        ["Total human effort", "Hrobot = Hsupervision + Htransfer + Hmaintenance"],
        ["Reclaimed effort", "Hreclaimed = H₀ − Hrobot"],
    ])
    p("p is the selected observed success fraction; hₜ and hᵣ are the human fractions of transfer/recovery (both 1 in the base case). Repeated attempts are a business assumption; they do not create a demonstrated autonomous recovery capability.")
    h("Worked base case", 2)
    table(["Step", "Calculated value"], [
        ["Attempt hours", f"630 × {m['mean_attempt_s']:.6f} / 3600 = {630*m['mean_attempt_s']/3600:.6f} h/month"],
        ["Monitoring", f"0.20 × attempt hours = {b['supervision_person_h_month']:.6f} person-h/month"],
        ["Transfer / recovery", f"630 × 180 / 3600 = {b['transfer_recovery_person_h_month']:.6f} person-h/month (p = 1)"],
        ["Maintenance", "10.000000 person-h/month"],
        ["Human effort with robot", f"{b['total_robot_related_person_h_month']:.6f} person-h/month"],
        ["Reclaimed effort", f"120 − {b['total_robot_related_person_h_month']:.6f} = {b['labour_reclaimed_person_h_month']:.6f} person-h/month"],
    ])
    p("Reclaimed time is a scenario opportunity to reallocate labour; it is not a measured reduction in headcount. The model counts all transfer labour even when the robot needs no active supervision.")

    # 6. Worked cost and power accounting.
    page("5 · Energy, operating cost and payback")
    p("Attempt energy estimate: Eₐ = (Wpositive/η + Eblower + Pelectronics × Tₐ) / 3600, in Wh. Per-success robot energy adds retries and electronics during recovery/transfer: Erobot = Eₐ/p + Pelectronics × (R + tₜ)/3600.")
    robot_kwh = a["sections_per_day"] * b["estimated_robot_Wh_per_successful_section"] * a["operating_days_per_month"] * 12 / 1000
    camera_kwh = b["estimated_system_kWh_year"] - robot_kwh
    table(["Energy estimate", "Calculated value"], [
        ["Robot energy / successful section", f"{b['estimated_robot_Wh_per_successful_section']:.6f} Wh"],
        ["Robot energy / modelled m²", f"{b['estimated_energy_Wh_per_modelled_m2']:.6f} Wh/m²; excludes external cameras"],
        ["Annual robot / external cameras", f"{robot_kwh:.6f} / {camera_kwh:.6f} kWh/year"],
        ["Annual system estimate", f"{b['estimated_system_kWh_year']:.6f} kWh/year"],
    ])
    h("Normalised annual operating-cost bridge", 2)
    p("κ = Hrobot/H₀ + Eannual × α + cspares + cconsumables. Capital is evaluated separately in payback.")
    table(["Cost component", "Fraction of annual manual baseline"], [
        ["Monitoring labour", f"{b['supervision_person_h_month']/120:.8f}"],
        ["Transfer / recovery labour", f"{b['transfer_recovery_person_h_month']/120:.8f}"],
        ["Maintenance labour", f"{10/120:.8f}"],
        ["Estimated energy", f"{b['energy_cost_ratio']:.8f}"],
        ["Battery / spares", "0.05000000"],
        ["Consumables", "0.03000000"],
        ["Total κ", f"{b['kappa']:.8f}"],
    ])
    p(f"NCRI = 1 − κ = {b['NCRI']:.8f} ({100*b['NCRI']:.2f}%). Annual savings ratio s = NCRI. Simple payback T = 12k/s = 12 × 0.65 / {b['NCRI']:.8f} = {b['payback_months']:.4f} months. Lifetime-normalised payback P = k/(L × s) = {b['lifetime_normalised_payback']:.6f}.")
    p("If s ≤ 0, payback is undefined; the script does not report an invented finite recovery time. Financing, tax, discounting and damage-related losses are outside this simple-payback calculation.")

    # 7. Queueing and existing source figure.
    page("6 · Capacity, queueing and labour graph")
    p("One server represents one robot-plus-transfer service process. λ = sections/day ÷ scheduled hours/day. μ = 3600/Tₛ. For M/M/1, ρ = λ/μ; Wq = λ/[μ(μ − λ)] and W = 1/(μ − λ), provided λ < μ.")
    table(["Queue output", "Base case"], [
        ["Service time Tₛ", f"{b['service_mean_s']:.6f} s = {b['service_mean_s']/60:.4f} min"],
        ["λ / μ", f"{b['arrival_rate_sections_h']:.6f} / {b['service_rate_sections_h']:.6f} sections/h"],
        ["Servers / utilisation", f"1 / {b['utilisation']:.6f}"],
        ["Mean queue wait Wq", f"{60*b['queue_wait_h']:.6f} min"],
        ["Mean queue + service W", f"{60*b['queue_system_h']:.6f} min"],
        ["Service-process hours/month", f"{b['robot_operating_h_month']:.6f} h, including assumed transfer"],
    ])
    picture(DATA / "queue_and_labour.png", "Existing final-analysis figure. Queue distributions are assumed. The M/D/1 curve is a reference, not a fitted arrival/service model. Labour includes monitoring, transfer/recovery and maintenance.")
    p("A mean queue wait can exceed one service duration because arrivals may find work already waiting. Queue delay is a stochastic scenario output, not an observed timing measurement from the robot film. If utilisation reaches 1, finite steady-state wait is not reported. Maintenance labour is costed but its detailed equipment-downtime schedule is not simulated.")

    # 8. Sensitivity with all eight parameters.
    page("7 · Sensitivity and break-even conditions")
    p("The script changes eight inputs one at a time to 70%, 100% and 130% of their base values. This is scenario sensitivity, not a financial confidence interval or a joint uncertainty model.")
    labels = {
        "supervision_fraction": "Supervision", "manual_person_hours_per_month": "Manual baseline hours",
        "capital_to_one_year_manual_cost_ratio": "Capital ratio", "between_section_transfer_s": "Transfer seconds",
        "failed_attempt_recovery_s": "Recovery seconds", "electricity_cost_per_kWh_to_one_year_manual_cost_ratio": "Energy-cost coefficient",
        "maintenance_person_hours_per_month": "Maintenance hours", "electromechanical_efficiency": "Drive efficiency",
    }
    rows = []
    for parameter, label in labels.items():
        part = sensitivity[sensitivity.parameter == parameter]
        lo, hi = part.iloc[0], part.iloc[2]
        rows.append([label, f"{100*lo.NCRI:.2f}% / {100*hi.NCRI:.2f}%", f"{lo.payback_months:.2f} / {hi.payback_months:.2f}"])
    table(["Input changed", "NCRI at −30% / +30%", "Payback months at −30% / +30%"], rows, size=8.5)
    picture(DATA / "business_sensitivity.png", "Existing final-analysis figure: supervision versus normalised operating-cost reduction and simple payback. Commercial assumptions are varied; savings are not field observations.")
    p(f"Base break-even supervision fraction: {b['break_even_supervision_fraction']:.6f} ({100*b['break_even_supervision_fraction']:.2f}%). Above this, annual operating savings become non-positive with other assumptions fixed. Maximum capital ratio for recovery within five assumed years: {b['maximum_capital_ratio_for_payback_within_life']:.6f}.")
    p("Recovery-time sensitivity is flat at p = 1 because no recovery is invoked in the base scenario. It matters when p < 1. Sensitivity holds other declared normalised inputs fixed; it is not a set of verified supplier-price forecasts.")

    # 9. Confidence scenario and limits.
    page("8 · Reliability uncertainty and interpretation")
    p("Thirty observed successes do not establish true success probability 1. The stored lower-confidence scenario substitutes the nominal proposed controller’s 95% Wilson lower bound for p, while retaining the same measured mean attempt/work inputs. This is a stress scenario, not a prediction of field failure rates.")
    table(["Output", "Observed-sample base", "Wilson-lower-bound scenario"], [
        ["Success probability used", "1.000000", f"{conservative['success_fraction']:.6f}"],
        ["Service capacity (sections/h)", f"{b['service_rate_sections_h']:.4f}", f"{conservative['service_rate_sections_h']:.4f}"],
        ["Monthly human effort (person-h)", f"{b['total_robot_related_person_h_month']:.4f}", f"{conservative['total_robot_related_person_h_month']:.4f}"],
        ["Monthly effort reclaimed (person-h)", f"{b['labour_reclaimed_person_h_month']:.4f}", f"{conservative['labour_reclaimed_person_h_month']:.4f}"],
        ["κ / NCRI", f"{b['kappa']:.5f} / {100*b['NCRI']:.2f}%", f"{conservative['kappa']:.5f} / {100*conservative['NCRI']:.2f}%"],
        ["Simple payback (months)", f"{b['payback_months']:.4f}", f"{conservative['payback_months']:.4f}"],
        ["Queue wait (min)", f"{60*b['queue_wait_h']:.4f}", f"{60*conservative['queue_wait_h']:.4f}"],
        ["Estimated system kWh/year", f"{b['estimated_system_kWh_year']:.4f}", f"{conservative['estimated_system_kWh_year']:.4f}"],
    ])
    h("Evidence boundary", 2)
    p("The frozen simulation varies initial pose, dust layout and incline under its declared protocol. This does not calibrate wet-glass friction, varying dirt colours, shadows, coating wear or hardware vacuum performance. Navigation uses simulator pose; rendered RGB drives contamination detection. Fine-dirt capture is a contact-gated reduced-order mass model, while the larger solid uses rigid-body pushing in a prescribed corridor.")
    p("Physical containment guards are part of the demonstrated mechanism; their success does not prove unguarded real-world traction. The dock’s charging and emptying counters are modelled service, not tested connectors, airflow or battery recovery.")
    h("Commercial evidence still needed", 2)
    p("Measure operator supervision/transfer time, site workload, actual electrical energy, component life, downtime and maintenance. Obtain compatible-module guidance, supplier quotations and pilot customer feedback. Keep these commercial uncertainties separate from controller Monte Carlo statistics.")

    # 10. Actual dossier FMEA.
    page("9 · Risk register and FMEA")
    p("RPN = severity × occurrence × detection, each scored on an ordinal 1–10 scale. These engineering judgements are not estimated probabilities or insurance loss rates. Scores below match the current technical dossier.")
    dossier = Document(ROOT / "submission/dossier_revision_v6/Group06_Dossier_Updated.docx")
    rows = []
    for row in dossier.tables[14].rows[1:]:
        cells = [c.text for c in row.cells]
        rows.append([cells[0], "/".join(cells[2:5]), cells[5], cells[7], cells[6]])
    table(["Failure mode", "S/O/D pre", "RPN pre", "RPN post", "Mitigation and evidence boundary"], rows, size=8.5, widths=[1.45, .6, .55, .55, 3.8])
    h("Commercial consequence", 2)
    p("Traction loss, jams and service failures can increase supervision, recovery, downtime and replacement costs. Those effects motivate the explicit labour terms and sensitivity cases. The current model does not assign a currency value or a probability to catastrophic panel damage.")
    p("No RPN improvement is claimed for coating scratches: hardware compatibility is untested. Other post-RPN reductions credit only the described software/simulation detection or final logged verification; they do not establish field certification.")

    # 11. Decision and proposed validation path.
    page("10 · Decision, deployment plan and viva notes")
    h("Decision supported by this model", 2)
    p(f"Under the base assumptions, recurring operating cost is {100*b['kappa']:.2f}% of the manual baseline, with simple capital recovery in about {b['payback_months']:.1f} months. The lower-confidence success scenario still produces positive operating savings, but real commercial feasibility depends on the declared labour/cost inputs and hardware performance.")
    table(["Next gate", "Evidence to collect before a deployment claim"], [
        ["Bench prototype", "Measured current/power, motor and brush loads, debris collection, structural attachment and module/coating compatibility."],
        ["Controlled inclined-panel trial", "Slip/stop behaviour, guarded containment, illumination challenges, edge cases and retained solid collection."],
        ["Operator pilot", "Measured supervision, transfers, failures, maintenance, workload and actual downtime; replace assumptions in inputs.json."],
        ["Commercial pilot", "Supplier quotes, customer acceptance, support effort and unit-economics evidence; compare against a documented manual baseline."],
    ])
    h("Short answers for the business viva", 2)
    table(["Question", "Answer"], [
        ["Why not quote a price?", "The course requires dimensionless comparisons; verified manufacturing costs and customer pricing are not available."],
        ["Where do the headline numbers come from?", "calculate(inputs, measured) uses retained full-mission timing/work plus labelled assumptions. Results are not typed in as target outputs."],
        ["Does this remove human work?", "No. The model charges monitoring, all inter-bay transfer/recovery and maintenance. Reclaimed hours represent reallocation potential."],
        ["Why not claim energy-yield restoration?", "Dust/contact clearance is not electrical PV output. No PV electrical model or field-yield study is part of this result."],
        ["Why does queueing matter?", "It tests whether scheduled workload is compatible with one system’s service capacity and exposes delay growth near saturation."],
    ])

    # 12. Reproduction and source mapping.
    page("11 · Files, reproduction and submission")
    table(["File / folder", "Purpose"], [
        ["analytics/final_business_model.py", "Actual executable business model, sensitivity and plots"],
        ["analytics/final_business_inputs.json", "Declared commercial/energy assumptions"],
        ["analytics/final_validation_v1/{summary.csv, protocol.json, statistics.json}", "Full 180-attempt summary, selection rules and confidence bound"],
        ["analytics/final_business_v1/{inputs.json, simulation_inputs.json, results.json}", "Retained input snapshot, measured calibration and final outputs"],
        ["analytics/final_business_v1/sensitivity.csv", "All 24 one-at-a-time cases (eight inputs × three factors)"],
        ["business_sensitivity.png / queue_and_labour.png", "Final-analysis figures embedded in this report"],
        ["analytics/final_validation_v1/tilt_*/...", "Bulk raw CSV/XZ telemetry in the full project repository, not duplicated in the business-only package"],
    ], size=8.5)
    h("Run the business model without rerunning MuJoCo", 2)
    p("From the group folder, use Python 3.12 with numpy, pandas and matplotlib (the existing simulation environment already includes them). The package includes requirements-business.txt for an independent business-only environment.")
    p("python analytics/final_business_model.py --validation-dir analytics/final_validation_v1 --output-dir analytics/business_recalculated")
    p("Use a new output directory: the script refuses to overwrite an existing analysis. It loads the saved final evaluation; no new robot simulation is required. Output files are results.json, inputs.json, simulation_inputs.json, sensitivity.csv and the two graphs.")
    h("What the rubric asks you to submit", 2)
    p("MS Teams: every member individually attaches the group’s final signed dossier PDF and clicks Turn in. GitHub faculty group folder: filled dossier Word/PDF in docs/, model XML, controller code, business script, genuine trial data and reproducible outputs. LinkedIn: final team-video URL in dossier Section 5. This business report is supporting documentation; it does not replace the Python business script.")
    p("Source: GRADING_RUBRIC_AND_EVALUATION_POLICY.md in sunny-nanade/NMIMS-MPSTME-MDRIIA-2026, checked at commit d5b7da279a034fcd08b090766ef96ed31b196684. The current 20-mark rubric supersedes the older milestone/optional IEEE-paper track; no PPT is required. See ASSIGNMENT_CHECKLIST.md in the package.")
    p("Commercial-input evidence is the declared assumption file, not an external quotation. Calculation evidence is the retained simulation summary and final business outputs. The report introduces no sourced Mumbai-yield percentage, measured commercial saving or fabricated market survey.")

    target = output / "Group06_Business_Model_Report.docx"
    document.save(target)
    sources = [DATA / name for name in ["results.json", "sensitivity.csv", "business_sensitivity.png", "queue_and_labour.png"]]
    sources += [ROOT / "analytics/final_business_model.py", ROOT / "analytics/final_business_inputs.json", ROOT / "analytics/final_validation_v1/summary.csv", ROOT / "analytics/final_validation_v1/protocol.json", ROOT / "analytics/final_validation_v1/statistics.json", ROOT / "submission/dossier_revision_v6/Group06_Dossier_Updated.docx"]
    manifest = {"sources": [{"path": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in sources], "word_sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "declared_document_sections": 12, "business_headlines": {key: b[key] for key in ["kappa", "NCRI", "payback_months", "labour_reclaimed_person_h_month"]}, "commercial_inputs": "Assumptions, not verified quotations or field savings", "faculty_requirement_source_commit": "d5b7da279a034fcd08b090766ef96ed31b196684"}
    (output / "REPORT_SOURCE_MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(target)
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT / "submission/business_model_report_v1")
    args = parser.parse_args()
    build(args.output_dir.resolve())
